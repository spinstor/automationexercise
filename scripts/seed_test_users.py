#!/usr/bin/env python
"""Seed the Excel test-data file used by the registration_from_excel suite.

Creates resources/data/test_users.xlsx (overwrites any existing file) with
the expected column headers and a few ready-to-use signup rows. Emails are
timestamped so each seeded file contains fresh, unique addresses that will
not clash with previous runs.

Example:
    python scripts/seed_test_users.py
    python scripts/seed_test_users.py --rows 5
"""

import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = [
    "name",
    "email",
    "password",
    "title",
    "birth_day",
    "birth_month",
    "birth_year",
    "first_name",
    "last_name",
    "company",
    "address1",
    "country",
    "state",
    "city",
    "zipcode",
    "mobile",
    "status",
    "result_message",
]

SAMPLE_USERS = [
    {
        "name": "Excel Alice",
        "password": "Passw0rd!123",
        "title": "Mr",
        "birth_day": 15,
        "birth_month": "June",
        "birth_year": 1991,
        "first_name": "Alice",
        "last_name": "Summers",
        "company": "Openpyxl QA",
        "address1": "12 Data Drive",
        "country": "India",
        "state": "Karnataka",
        "city": "Bengaluru",
        "zipcode": "560001",
        "mobile": "9988776655",
        "status": "",
        "result_message": "",
    },
    {
        "name": "Excel Bob",
        "password": "Passw0rd!123",
        "title": "Mr",
        "birth_day": 22,
        "birth_month": "March",
        "birth_year": 1988,
        "first_name": "Bob",
        "last_name": "Fields",
        "company": "Sheetware Ltd",
        "address1": "8 Row Road",
        "country": "United States",
        "state": "California",
        "city": "San Francisco",
        "zipcode": "94105",
        "mobile": "4155550101",
        "status": "",
        "result_message": "",
    },
    {
        "name": "Excel Carol",
        "password": "Passw0rd!123",
        "title": "Mrs",
        "birth_day": 3,
        "birth_month": "November",
        "birth_year": 1995,
        "first_name": "Carol",
        "last_name": "Denver",
        "company": "Cellrange Inc",
        "address1": "77 Workbook Way",
        "country": "India",
        "state": "Tamil Nadu",
        "city": "Chennai",
        "zipcode": "600002",
        "mobile": "9123456780",
        "status": "",
        "result_message": "",
    },
]


def build(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(len(header) + 2, 14)

    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    emails = [f"excel.user.{timestamp}.{i}@example.com" for i in range(rows)]

    for idx, sample in zip(range(rows), [SAMPLE_USERS[i % len(SAMPLE_USERS)] for i in range(rows)]):
        row_data = sample.copy()
        row_data["email"] = emails[idx]
        for col_idx, header in enumerate(COLUMNS, start=1):
            ws.cell(row=idx + 2, column=col_idx, value=row_data.get(header, ""))

    out = Path("resources/data/test_users.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.close()
    print("Seeded %d user rows -> %s" % (rows, out))
    for email in emails:
        print("  - %s" % email)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, default=3, help="number of signup rows to generate")
    args = parser.parse_args()
    build(args.rows)


if __name__ == "__main__":
    main()